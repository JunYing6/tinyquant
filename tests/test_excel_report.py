"""Excel report export tests."""

from __future__ import annotations

from datetime import datetime, timezone
from typing import Any

import pytest
from openpyxl import load_workbook

from engines.fast import FastBacktestEngine
from tools.data import Bar, DataRequest, InMemoryGateway, Session, TradingPhase
from tools.excel_report import (
    _localized_stats,
    default_output_dir,
    export_backtest_excel,
)
from trading_nodes_base.methods import BaseTimeSelection
from trading_nodes_base.strategies import BaseStrategy
from trading_nodes_base.types import ExecutionMode, ExecutionRequest
from trading_nodes_base.factors import KlineTimingFactor, TickTimingFactor


def _instant(day: str, clock: str = "15:00:00") -> datetime:
    return datetime.strptime(f"{day} {clock}", "%Y%m%d %H:%M:%S").replace(tzinfo=timezone.utc)


def _bar(day: str, price: float, code: str = "000001.SZ") -> Bar:
    instant = _instant(day)
    return Bar(schema_version="1", event_id=None, instrument_id=code, asset_type="equity", effective_time=instant, event_time=instant, available_at=instant, trading_date=instant.date(), source="test", quality="valid", metadata={}, frequency="1d", interval_start=instant, interval_end=instant, open=price, high=price, low=price, close=price, volume=0, turnover=0, is_complete=True, price_basis="raw")


def _session(day: str) -> Session:
    phase = TradingPhase(name="regular", start=_instant(day, "09:00:00"), end=_instant(day, "15:00:00"), accepts_trades=True, accepts_quotes=True)
    return Session(market="CN", trading_date=_instant(day).date(), timezone="UTC", phases=(phase,))


def _gateway(bars: list[Bar]) -> InMemoryGateway:
    days = ["20240102", "20240103"]
    return InMemoryGateway(bars=bars, events=[], sessions=[_session(day) for day in days])


class PassiveQueryFactor(KlineTimingFactor):
    def __init__(self) -> None:
        super().__init__("passive-query")

    def get_query_lst(self, date: object, codes: list[str] | None = None) -> list[DataRequest]:
        self._data_clear()
        self.sign["fit"] = True
        return []

    def on_bar(self, bar: object) -> list:
        return []


class PassiveIntentExecutor(TickTimingFactor):
    execution_role = "intent_executor"
    accepted_intent_actions = frozenset({"BUY"})

    def __init__(self) -> None:
        super().__init__("passive-executor")

    def get_query_lst(self, date: object, codes: list[str] | None = None) -> list[DataRequest]:
        self._data_clear()
        self.sign["fit"] = True
        return []


class BuyingFastStrategy(BaseStrategy):
    supports_fast_backtest = True

    def __init__(self) -> None:
        timer = BaseTimeSelection("buy-timer", [PassiveQueryFactor()], [PassiveIntentExecutor()])
        super().__init__("buy-fast", timer=timer)
        self._queued_once = False

    def _run_daily_pipeline(self) -> None:
        super()._run_daily_pipeline()
        if not self._queued_once:
            self._pending_orders.append(
                ExecutionRequest("000001.SZ", "BUY", "15:00:00", price=0, volume=100, mode=ExecutionMode.MARKET)
            )
            self._queued_once = True


def _run_engine() -> FastBacktestEngine:
    bars = [_bar("20240102", 10.0), _bar("20240103", 11.0)]
    engine = FastBacktestEngine(BuyingFastStrategy(), "20240102", "20240103", initial_capital=100_000, mode="fast", data_gateway=_gateway(bars), progress_bar=False)
    engine.run()
    return engine


def test_export_backtest_excel_writes_workbook(tmp_path) -> None:
    engine = _run_engine()

    path = export_backtest_excel(engine, output_dir=tmp_path)

    assert path.exists()
    assert path.parent == tmp_path
    assert path.name.startswith("buy-fast_20240102_20240103_")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["策略概览", "权益曲线", "交易记录", "持仓明细", "月度收益"]
    overview = workbook["策略概览"]
    assert overview["A1"].value == "tinyquant 回测报告"
    assert overview["B4"].value == "buy-fast"
    equity = workbook["权益曲线"]
    assert equity["A2"].value == "20240102"
    assert equity["B3"].value == pytest.approx(engine.equity_curve[-1]["equity"], abs=0.01)
    positions = workbook["持仓明细"]
    assert positions["B2"].value == "000001.SZ"
    assert positions["D2"].value == pytest.approx(10.0, abs=0.01)


def test_export_uses_default_output_dir(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("TINYQUANT_EXCEL_DIR", str(tmp_path))
    assert default_output_dir() == tmp_path

    engine = _run_engine()
    path = export_backtest_excel(engine)

    assert path.parent == tmp_path
    assert path.exists()


def test_localized_stats_scales_percent_keys() -> None:
    stats = {"total_return": 0.12, "sharpe": 1.5, "win_rate": 0.6, "trading_days": 2, "final_equity": 112000.0}

    localized = _localized_stats(stats)

    assert localized["总收益率 (%)"] == pytest.approx(12.0)
    assert localized["胜率 (%)"] == pytest.approx(60.0)
    assert localized["夏普比率"] == pytest.approx(1.5)
    assert localized["交易天数"] == 2
    assert localized["最终权益"] == pytest.approx(112000.0)
