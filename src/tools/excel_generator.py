"""Excel backtest report generator.

Output layout: one multi-sheet workbook holding the strategy overview,
equity curve, trade log, daily positions and monthly returns.
"""

from __future__ import annotations

import os
from datetime import datetime
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

TradeRecord = tuple[str, Any, str, float, int]


class ExcelReportGenerator:
    """Render backtest records as a styled multi-sheet Excel workbook.

    Args:
        equity_curve: daily equity rows (``date``, ``equity``, ``balance``, ...).
        daily_positions: per-day holdings rows (``date``, ``positions``,
            ``close_prices``).
        trade_log: trades grouped by date: ``{date: [(time, code, price, qty)]}``.
        stats_dict: performance metrics keyed by the report's Chinese labels.
        strategy_name: strategy name shown in the overview sheet.
        start_date: backtest start date (YYYYMMDD).
        end_date: backtest end date (YYYYMMDD).
        initial_capital: starting capital of the backtest.
    """

    HEADER_FONT = Font(name='微软雅黑', bold=True, color='FFFFFF')
    HEADER_FILL = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    TITLE_FONT = Font(name='微软雅黑', size=14, bold=True)
    LABEL_FONT = Font(name='微软雅黑', bold=True)
    THIN_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    NUM_FORMAT = '#,##0.00'

    def __init__(
        self,
        equity_curve: list[dict[str, Any]],
        daily_positions: list[dict[str, Any]],
        trade_log: dict[str, list[tuple[Any, str, float, int]]],
        stats_dict: dict[str, Any],
        strategy_name: str,
        start_date: str,
        end_date: str,
        initial_capital: float,
    ) -> None:
        self.equity_curve = equity_curve
        self.daily_positions = daily_positions
        self.trade_log = trade_log
        self.stats_dict = stats_dict
        self.strategy_name = strategy_name
        self.start_date = start_date
        self.end_date = end_date
        self.initial_capital = initial_capital
        self.member_count = 0

    def generate(self, save_path: str) -> str:
        """Write the workbook to ``save_path`` and return the path."""
        os.makedirs(os.path.dirname(save_path) or '.', exist_ok=True)

        wb = Workbook()
        self._create_sheets(wb)
        if 'Sheet' in wb.sheetnames:
            del wb['Sheet']

        wb.save(save_path)
        return save_path

    def _report_title(self) -> str:
        return 'tinyquant 回测报告'

    def _create_sheets(self, wb: Workbook) -> None:
        self._create_overview_sheet(wb)
        self._create_equity_sheet(wb)
        self._create_trades_sheet(wb)
        self._create_positions_sheet(wb)
        self._create_monthly_sheet(wb)

    def _set_header_row(self, ws: Worksheet, row: int, headers: list[str]) -> None:
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=header)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = Alignment(horizontal='center')
            cell.border = self.THIN_BORDER

    def _set_col_widths(self, ws: Worksheet, widths: dict[str, int]) -> None:
        for col_letter, width in widths.items():
            ws.column_dimensions[col_letter].width = width

    def _create_overview_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet('策略概览', 0)

        ws.merge_cells('A1:F1')
        ws['A1'] = self._report_title()
        ws['A1'].font = self.TITLE_FONT
        ws['A1'].alignment = Alignment(horizontal='center')

        ws['A2'] = '生成时间'
        ws['B2'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        ws['A2'].font = self.LABEL_FONT

        info_rows = [
            ['策略名称', self.strategy_name, '回测期间', f'{self.start_date} - {self.end_date}'],
            ['初始资金', self.initial_capital, '交易天数', self.stats_dict.get('交易天数', 0)],
        ]
        if self.member_count:
            info_rows.append(['成员策略数', self.member_count, '', ''])
        for r_idx, row_data in enumerate(info_rows, start=4):
            for c_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                if c_idx in [1, 3]:
                    cell.font = self.LABEL_FONT
                if isinstance(value, float):
                    cell.number_format = self.NUM_FORMAT

        ws.cell(row=7, column=1, value='绩效指标').font = self.LABEL_FONT
        metric_headers = ['指标', '数值', '指标', '数值']
        self._set_header_row(ws, 8, metric_headers)

        metrics = self._get_metrics_data()
        for r_idx, row_data in enumerate(metrics, start=9):
            for c_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                if c_idx in [1, 3]:
                    cell.font = self.LABEL_FONT
                if isinstance(value, (int, float)):
                    cell.number_format = self.NUM_FORMAT
                cell.border = self.THIN_BORDER
                cell.alignment = Alignment(horizontal='center')

        self._set_col_widths(ws, {'A': 16, 'B': 20, 'C': 16, 'D': 20, 'E': 5, 'F': 5})

    def _get_metrics_data(self) -> list[tuple[str, Any, str, Any]]:
        flat_trades = self._flatten_trades()
        return [
            (
                '总收益率',
                f"{self.stats_dict.get('总收益率 (%)', 0):.2f}%",
                '年化收益率',
                f"{self.stats_dict.get('年化收益率 (%)', 0):.2f}%"
            ),
            (
                '最大回撤',
                f"{self.stats_dict.get('最大回撤 (%)', 0):.2f}%",
                '年化波动率',
                f"{self.stats_dict.get('年化波动率 (%)', 0):.2f}%"
            ),
            (
                '夏普比率',
                f"{self.stats_dict.get('夏普比率', 0):.2f}",
                '胜率',
                f"{self.stats_dict.get('胜率 (%)', 0):.2f}%"
            ),
            (
                '最终权益',
                self.stats_dict.get('最终权益', 0),
                '总交易次数',
                self.stats_dict.get('总交易次数', len(flat_trades))
            ),
        ]

    def _create_equity_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet('权益曲线')

        headers = ['日期', '权益', '现金', '持仓市值', '持仓数量', '日收益率(%)', '累计收益率(%)', '回撤(%)']
        self._set_header_row(ws, 1, headers)

        df = pd.DataFrame(self.equity_curve)
        if df.empty:
            return

        df['daily_return'] = df['equity'].pct_change().fillna(0)
        df['cum_return'] = df['equity'] / self.initial_capital - 1
        df['drawdown'] = df['equity'] / df['equity'].cummax() - 1

        for row_idx, row in df.iterrows():
            r = row_idx + 2
            ws.cell(row=r, column=1, value=row['date']).border = self.THIN_BORDER
            ws.cell(row=r, column=2, value=round(row['equity'], 2)).border = self.THIN_BORDER
            ws.cell(row=r, column=3, value=round(row['balance'], 2)).border = self.THIN_BORDER
            ws.cell(row=r, column=4, value=round(row.get('position_value', 0), 2)).border = self.THIN_BORDER
            ws.cell(row=r, column=5, value=int(row.get('position_count', 0))).border = self.THIN_BORDER
            ws.cell(row=r, column=6, value=round(row['daily_return'] * 100, 4)).border = self.THIN_BORDER
            ws.cell(row=r, column=7, value=round(row['cum_return'] * 100, 2)).border = self.THIN_BORDER
            ws.cell(row=r, column=8, value=round(row['drawdown'] * 100, 2)).border = self.THIN_BORDER

            for col in [2, 3, 4]:
                ws.cell(row=r, column=col).number_format = self.NUM_FORMAT

        last_row = len(df) + 2
        ws.cell(row=last_row, column=1, value='合计/均值').font = self.LABEL_FONT
        ws.cell(row=last_row, column=2, value=round(df['equity'].iloc[-1], 2)).number_format = self.NUM_FORMAT
        ws.cell(row=last_row, column=6, value=round(df['daily_return'].mean() * 100, 4))
        ws.cell(row=last_row, column=8, value=round(df['drawdown'].min() * 100, 2))

        self._add_equity_chart(ws, len(df) + 1, last_row)
        self._set_col_widths(ws, {c: 15 for c in 'ABCDEFGH'})

    def _add_equity_chart(self, ws: Worksheet, data_rows: int, start_row: int) -> None:
        chart = LineChart()
        chart.title = "权益曲线"
        chart.y_axis.title = "权益"
        chart.x_axis.title = "日期"
        chart.style = 10
        chart.width = 30
        chart.height = 15

        data = Reference(ws, min_col=2, min_row=1, max_row=data_rows)
        cats = Reference(ws, min_col=1, min_row=2, max_row=data_rows)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.series[0].graphicalProperties.line.width = 25000

        ws.add_chart(chart, f"A{start_row + 2}")

    def _create_trades_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet('交易记录')

        headers = ['日期', '时间', '股票代码', '方向', '价格', '数量', '金额']
        self._set_header_row(ws, 1, headers)

        flat_trades = self._flatten_trades()
        for r_idx, trade in enumerate(flat_trades, start=2):
            date_str, time_val, code, price, num = trade
            direction = '买入' if num > 0 else '卖出'
            amount = abs(num) * price

            ws.cell(row=r_idx, column=1, value=date_str).border = self.THIN_BORDER
            ws.cell(row=r_idx, column=2, value=str(time_val)).border = self.THIN_BORDER
            ws.cell(row=r_idx, column=3, value=code).border = self.THIN_BORDER

            dir_cell = ws.cell(row=r_idx, column=4, value=direction)
            dir_cell.border = self.THIN_BORDER
            dir_cell.font = Font(color='FF0000') if num > 0 else Font(color='008000')

            ws.cell(row=r_idx, column=5, value=round(price, 2)).border = self.THIN_BORDER
            ws.cell(row=r_idx, column=6, value=abs(int(num))).border = self.THIN_BORDER
            ws.cell(row=r_idx, column=7, value=round(amount, 2)).border = self.THIN_BORDER

            ws.cell(row=r_idx, column=5).number_format = self.NUM_FORMAT
            ws.cell(row=r_idx, column=7).number_format = self.NUM_FORMAT

        if flat_trades:
            self._add_trade_summary(ws, flat_trades)

        self._set_col_widths(ws, {c: 15 for c in 'ABCDEFG'})

    def _add_trade_summary(self, ws: Worksheet, flat_trades: list[TradeRecord]) -> None:
        summary_row = len(flat_trades) + 3
        ws.cell(row=summary_row, column=1, value='交易统计').font = self.LABEL_FONT

        buys = [t for t in flat_trades if t[4] > 0]
        sells = [t for t in flat_trades if t[4] < 0]

        ws.cell(row=summary_row + 1, column=1, value='总买入次数')
        ws.cell(row=summary_row + 1, column=2, value=len(buys))
        ws.cell(row=summary_row + 2, column=1, value='总卖出次数')
        ws.cell(row=summary_row + 2, column=2, value=len(sells))

        total_amount = sum(abs(t[4]) * t[3] for t in flat_trades)
        ws.cell(row=summary_row + 3, column=1, value='总交易金额')
        ws.cell(row=summary_row + 3, column=2, value=round(total_amount, 2))
        ws.cell(row=summary_row + 3, column=2).number_format = self.NUM_FORMAT

    def _create_positions_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet('持仓明细')

        headers = ['日期', '股票代码', '持仓数量', '收盘价', '市值']
        self._set_header_row(ws, 1, headers)

        r_idx = 2
        for pos_data in self.daily_positions:
            date_str = pos_data['date']
            positions = pos_data.get('positions', {})
            close_prices = pos_data.get('close_prices', {})

            if not positions:
                ws.cell(row=r_idx, column=1, value=date_str).border = self.THIN_BORDER
                ws.cell(row=r_idx, column=2, value='空仓').border = self.THIN_BORDER
                for c in range(3, 6):
                    ws.cell(row=r_idx, column=c, value=0).border = self.THIN_BORDER
                r_idx += 1
                continue

            for code, quantity in positions.items():
                price = close_prices.get(code, 0)
                market_value = quantity * price

                ws.cell(row=r_idx, column=1, value=date_str).border = self.THIN_BORDER
                ws.cell(row=r_idx, column=2, value=code).border = self.THIN_BORDER
                ws.cell(row=r_idx, column=3, value=quantity).border = self.THIN_BORDER
                ws.cell(row=r_idx, column=4, value=round(price, 2)).border = self.THIN_BORDER
                ws.cell(row=r_idx, column=5, value=round(market_value, 2)).border = self.THIN_BORDER

                ws.cell(row=r_idx, column=4).number_format = self.NUM_FORMAT
                ws.cell(row=r_idx, column=5).number_format = self.NUM_FORMAT
                r_idx += 1

        self._set_col_widths(ws, {c: 15 for c in 'ABCDE'})

    def _create_monthly_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet('月度收益')

        headers = ['月份', '月收益率(%)', '累计收益率(%)']
        self._set_header_row(ws, 1, headers)

        df = pd.DataFrame(self.equity_curve)
        if df.empty:
            return

        df['date'] = pd.to_datetime(df['date'], format='%Y%m%d')
        df['month'] = df['date'].dt.to_period('M')

        monthly = df.groupby('month').agg(
            start_equity=('equity', 'first'),
            end_equity=('equity', 'last')
        ).reset_index()

        monthly['return_pct'] = (monthly['end_equity'] / monthly['start_equity'] - 1) * 100
        monthly['cum_return_pct'] = (monthly['end_equity'] / self.initial_capital - 1) * 100

        for r_idx, row in monthly.iterrows():
            r = r_idx + 2
            ws.cell(row=r, column=1, value=str(row['month'])).border = self.THIN_BORDER

            ret_cell = ws.cell(row=r, column=2, value=round(row['return_pct'], 2))
            ret_cell.border = self.THIN_BORDER
            if row['return_pct'] > 0:
                ret_cell.font = Font(color='FF0000')
            elif row['return_pct'] < 0:
                ret_cell.font = Font(color='008000')

            ws.cell(row=r, column=3, value=round(row['cum_return_pct'], 2)).border = self.THIN_BORDER

        if len(monthly) > 1:
            self._add_monthly_chart(ws, len(monthly))

        self._set_col_widths(ws, {c: 18 for c in 'ABC'})

    def _add_monthly_chart(self, ws: Worksheet, month_count: int) -> None:
        chart = LineChart()
        chart.title = "月度收益率"
        chart.y_axis.title = "收益率 (%)"
        chart.style = 10
        chart.width = 25
        chart.height = 12

        data = Reference(ws, min_col=2, min_row=1, max_row=month_count + 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=month_count + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)

        ws.add_chart(chart, f"A{month_count + 4}")

    def _flatten_trades(self) -> list[TradeRecord]:
        flat: list[TradeRecord] = []
        for date_str in sorted(self.trade_log.keys()):
            for trade in self.trade_log[date_str]:
                time_val, code, price, num = trade
                flat.append((date_str, time_val, code, price, num))
        return flat


class StreamExcelReportGenerator(ExcelReportGenerator):
    def __init__(
        self,
        equity_curve: list[dict[str, Any]],
        daily_positions: list[dict[str, Any]],
        trade_log: dict[str, list[tuple[Any, str, float, int]]],
        stats_dict: dict[str, Any],
        strategy_name: str,
        start_date: str,
        end_date: str,
        initial_capital: float,
        member_curve: list[dict[str, Any]],
        member_trade_logs: dict[str, dict],
    ) -> None:
        super().__init__(equity_curve, daily_positions, trade_log, stats_dict, strategy_name, start_date, end_date, initial_capital)
        self.member_curve = member_curve
        self.member_trade_logs = member_trade_logs
        self.member_count = len(self._member_names())

    def _report_title(self) -> str:
        return 'tinyquant 组合回测报告'

    def _create_sheets(self, wb: Workbook) -> None:
        super()._create_sheets(wb)
        self._create_member_summary_sheet(wb)
        self._create_member_equity_sheet(wb)
        self._create_member_weight_sheet(wb)

    def _member_names(self) -> list[str]:
        if self.member_curve:
            first = self.member_curve[0]
            if first.get('members'):
                return list(first['members'].keys())
            if first.get('weights'):
                return list(first['weights'].keys())
        return list(self.member_trade_logs.keys())

    def _member_trade_count(self, name: str) -> int:
        log = self.member_trade_logs.get(name, {})
        return sum(len(trades) for trades in log.values())

    def _create_member_summary_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet('成员策略汇总')
        headers = ['成员', '初始权重', '最新权重', '最终权益', '总收益(%)', '交易次数', '持仓市值']
        self._set_header_row(ws, 1, headers)

        names = self._member_names()
        first = self.member_curve[0] if self.member_curve else {}
        last = self.member_curve[-1] if self.member_curve else {}

        for r_idx, name in enumerate(names, start=2):
            init_w = (first.get('weights') or {}).get(name, 0.0)
            last_w = (last.get('weights') or {}).get(name, 0.0)
            init_eq = (first.get('members') or {}).get(name, {}).get('equity', 0.0)
            last_eq = (last.get('members') or {}).get(name, {}).get('equity', 0.0)
            ret = (last_eq / init_eq - 1) * 100 if init_eq else 0.0
            pv = (last.get('members') or {}).get(name, {}).get('position_value', 0.0)
            row = [name, round(init_w, 4), round(last_w, 4), round(last_eq, 2), round(ret, 2), self._member_trade_count(name), round(pv, 2)]
            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                cell.border = self.THIN_BORDER
                cell.alignment = Alignment(horizontal='center')
                if isinstance(value, float):
                    cell.number_format = self.NUM_FORMAT

        self._set_col_widths(ws, {col: 14 for col in 'ABCDEFG'})

    def _create_member_equity_sheet(self, wb: Workbook) -> None:
        self._create_member_curve_sheet(wb, '成员权益曲线', 'members', 'equity', '成员权益曲线')

    def _create_member_weight_sheet(self, wb: Workbook) -> None:
        self._create_member_curve_sheet(wb, '权重演变', 'weights', None, '权重演变')

    def _create_member_curve_sheet(self, wb: Workbook, sheet_name: str, key: str, value_key: str | None, title: str) -> None:
        ws = wb.create_sheet(sheet_name)
        names = self._member_names()
        headers = ['日期'] + names
        self._set_header_row(ws, 1, headers)

        rows = len(self.member_curve)
        for r_idx, snap in enumerate(self.member_curve, start=2):
            date_cell = ws.cell(row=r_idx, column=1, value=snap.get('date'))
            date_cell.border = self.THIN_BORDER
            data = snap.get(key) or {}
            for c_idx, name in enumerate(names, start=2):
                if value_key:
                    value = (data.get(name) or {}).get(value_key, 0.0)
                else:
                    value = data.get(name, 0.0)
                cell = ws.cell(row=r_idx, column=c_idx, value=round(value, 4))
                cell.border = self.THIN_BORDER
                if isinstance(value, float):
                    cell.number_format = self.NUM_FORMAT

        if rows:
            chart = LineChart()
            chart.title = title
            chart.y_axis.title = title
            chart.x_axis.title = "日期"
            chart.style = 10
            chart.width = 30
            chart.height = 15
            data = Reference(ws, min_col=2, min_row=1, max_row=rows + 1, max_col=len(names) + 1)
            cats = Reference(ws, min_col=1, min_row=2, max_row=rows + 1)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            ws.add_chart(chart, f"A{rows + 3}")

        self._set_col_widths(ws, {chr(ord('A') + i): 15 for i in range(len(names) + 1)})
